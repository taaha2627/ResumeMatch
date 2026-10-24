
import io
import logging
import os
import re
from collections import Counter
from typing import Optional

import fitz
import pytesseract
from PIL import Image
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from job_search import router as job_search_router


# =========================================================
# 1. APPLICATION CONFIGURATION
# =========================================================

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("resume_matcher")

APP_NAME = "ResumeMatch AI Career Studio"
MAX_FILE_SIZE = 10 * 1024 * 1024

TESSERACT_PATH = os.getenv(
    "TESSERACT_CMD",
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
)

if os.path.isfile(TESSERACT_PATH):
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH

app = FastAPI(
    title=APP_NAME,
    description="Resume analysis, skill extraction, and job matching.",
    version="2.0.0",
)

# Add your deployed frontend URL to this list when deploying.
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Register routers only after creating the FastAPI app.
app.include_router(job_search_router)


# =========================================================
# 2. SKILL DATABASE
# =========================================================

SKILL_GROUPS = {
    "Programming": [
        "python", "java", "javascript", "typescript", "c++", "c#",
        "c", "php", "ruby", "go", "rust", "kotlin", "swift",
    ],
    "Web Development": [
        "html", "css", "react", "angular", "vue.js", "node.js",
        "express.js", "fastapi", "flask", "django", "spring boot",
        "next.js", "rest api", "graphql",
    ],
    "AI and Machine Learning": [
        "machine learning", "deep learning", "artificial intelligence",
        "neural networks", "supervised learning", "unsupervised learning",
        "reinforcement learning", "natural language processing",
        "nlp", "computer vision", "generative ai", "llm",
        "large language models", "rag", "feature engineering",
        "model training", "model evaluation", "scikit-learn",
        "tensorflow", "pytorch", "keras", "opencv",
    ],
    "Data Science": [
        "numpy", "pandas", "matplotlib", "seaborn", "plotly",
        "data analysis", "data visualization", "statistics",
        "power bi", "tableau", "excel", "data mining",
    ],
    "Databases": [
        "sql", "mysql", "postgresql", "mongodb", "sqlite",
        "oracle database", "redis", "database management",
    ],
    "Cloud and DevOps": [
        "aws", "amazon web services", "azure", "microsoft azure",
        "google cloud", "docker", "kubernetes", "jenkins",
        "git", "github", "gitlab", "ci/cd", "devops",
        "terraform", "linux", "nginx",
    ],
    "Cybersecurity and Networking": [
        "cybersecurity", "network security", "cryptography",
        "blockchain", "ethical hacking", "computer networks",
    ],
    "Core Concepts": [
        "data structures", "algorithms", "object-oriented programming",
        "oops", "operating systems", "computer networks",
        "software engineering", "agile", "problem solving",
    ],
}

ALL_SKILLS = sorted(
    {skill for group in SKILL_GROUPS.values() for skill in group},
    key=len,
    reverse=True,
)

ROLE_SKILLS = {
    "AI/ML Engineer": [
        "python", "machine learning", "deep learning",
        "scikit-learn", "pytorch", "tensorflow", "statistics",
    ],
    "Computer Vision Engineer": [
        "python", "computer vision", "opencv",
        "deep learning", "pytorch", "tensorflow",
    ],
    "Data Scientist": [
        "python", "machine learning", "statistics",
        "pandas", "numpy", "sql", "data visualization",
    ],
    "Data Analyst": [
        "sql", "excel", "python", "pandas",
        "power bi", "tableau", "data visualization",
    ],
    "Full Stack Developer": [
        "html", "css", "javascript", "react",
        "node.js", "express.js", "sql", "git",
    ],
    "Backend Developer": [
        "python", "java", "sql", "rest api",
        "fastapi", "django", "node.js", "docker",
    ],
    "Software Engineer": [
        "python", "java", "data structures",
        "algorithms", "git", "sql", "object-oriented programming",
    ],
}


# =========================================================
# 3. TEXT NORMALIZATION AND EXTRACTION
# =========================================================

def normalize_text(text: str) -> str:
    """Normalize whitespace without deleting useful punctuation."""
    text = (text or "").replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract embedded PDF text and OCR pages that appear scanned."""
    extracted_pages = []

    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
            for page in pdf:
                page_text = page.get_text("text").strip()

                # OCR is used when a page contains little embedded text.
                if len(page_text) < 30:
                    try:
                        pixmap = page.get_pixmap(
                            matrix=fitz.Matrix(2, 2),
                            alpha=False,
                        )
                        image = Image.open(
                            io.BytesIO(pixmap.tobytes("png"))
                        )
                        ocr_text = pytesseract.image_to_string(image)
                        if len(ocr_text.strip()) > len(page_text):
                            page_text = ocr_text
                    except Exception as exc:
                        logger.warning("OCR failed for one PDF page: %s", exc)

                extracted_pages.append(page_text)

    except Exception as exc:
        logger.exception("PDF extraction failed")
        raise HTTPException(
            status_code=422,
            detail="Unable to read this PDF. Check that the file is valid.",
        ) from exc

    return normalize_text("\n".join(extracted_pages))


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract paragraphs and table contents from a DOCX document."""
    try:
        document = Document(io.BytesIO(file_bytes))
        parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if cells:
                    parts.append(" | ".join(cells))

        return normalize_text("\n".join(parts))

    except Exception as exc:
        logger.exception("DOCX extraction failed")
        raise HTTPException(
            status_code=422,
            detail="Unable to read this DOCX document. Check that the file is valid.",
        ) from exc


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    extension = os.path.splitext(filename.lower())[1]

    if extension == ".pdf":
        return extract_pdf_text(file_bytes)

    if extension == ".docx":
        return extract_docx_text(file_bytes)

    raise HTTPException(
        status_code=400,
        detail="Unsupported file format. Upload a PDF or DOCX file.",
    )


# =========================================================
# 4. SKILL DETECTION
# =========================================================

def contains_skill(text: str, skill: str) -> bool:
    """Match a skill as a phrase, not merely as part of another word."""
    pattern = r"(?<!\w)" + re.escape(skill.lower()) + r"(?!\w)"
    return re.search(pattern, text.lower()) is not None


def extract_skills(text: str) -> list[str]:
    normalized = normalize_text(text).lower()
    found = []

    for skill in ALL_SKILLS:
        if contains_skill(normalized, skill):
            found.append(skill)

    # Stable order with no duplicates.
    return sorted(set(found))


# =========================================================
# 5. RESUME SECTION DETECTION
# =========================================================

SECTION_PATTERNS = {
    "contact": [
        r"\bemail\b", r"\bphone\b", r"\bcontact\b",
        r"\blinkedin\b", r"\bgit(hub)?\b",
    ],
    "education": [
        r"\beducation\b", r"\bacademic background\b",
        r"\bqualification(s)?\b",
    ],
    "skills": [
        r"\btechnical skills\b", r"\bskills\b",
        r"\btechnologies\b", r"\btechnical expertise\b",
    ],
    "projects": [
        r"\bprojects\b", r"\bacademic projects\b",
        r"\bpersonal projects\b",
    ],
    "experience": [
        r"\bexperience\b", r"\binternship(s)?\b",
        r"\bwork history\b", r"\bprofessional experience\b",
    ],
    "summary": [
        r"\bsummary\b", r"\bcareer objective\b",
        r"\bprofessional profile\b", r"\bobjective\b",
    ],
    "certifications": [
        r"\bcertifications?\b", r"\blicenses\b",
        r"\bcertificates\b",
    ],
    "achievements": [
        r"\bachievements?\b", r"\bawards\b",
        r"\baccomplishments\b", r"\bcompetitions\b",
    ],
}


def detect_sections(text: str) -> dict[str, bool]:
    normalized = normalize_text(text).lower()
    sections = {}

    for section, patterns in SECTION_PATTERNS.items():
        sections[section] = any(
            re.search(pattern, normalized)
            for pattern in patterns
        )

    # Contact information may exist without a "Contact" heading.
    has_email = re.search(
        r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b",
        text,
        flags=re.IGNORECASE,
    )
    has_phone = re.search(r"(?<!\w)(?:\+?\d[\d ()-]{7,}\d)(?!\w)", text)

    sections["contact"] = bool(
        sections["contact"] or has_email or has_phone
    )
    return sections


# =========================================================
# 6. RESUME QUALITY SCORING
# =========================================================

def calculate_resume_quality(
    resume_text: str,
    resume_skills: Optional[list[str]] = None,
) -> dict:
    """
    Calculate a heuristic resume-quality score out of 100.
    This is not an AI hiring prediction.
    """
    text = normalize_text(resume_text)
    lower = text.lower()
    words = re.findall(r"\b[\w+#.-]+\b", text)
    word_count = len(words)
    skills = resume_skills if resume_skills is not None else extract_skills(text)
    sections = detect_sections(text)

    checks = {
        "contact_information": sections["contact"],
        "education_section": sections["education"],
        "skills_section": sections["skills"],
        "projects_section": sections["projects"],
        "experience_or_internship": sections["experience"],
        "summary_or_objective": sections["summary"],
        "certifications": sections["certifications"],
        "achievements": sections["achievements"],
    }

    # Structural score: 40 points.
    structural_weights = {
        "contact_information": 8,
        "education_section": 7,
        "skills_section": 7,
        "projects_section": 6,
        "experience_or_internship": 5,
        "summary_or_objective": 4,
        "certifications": 1,
        "achievements": 2,
    }

    structural_score = sum(
        weight for key, weight in structural_weights.items()
        if checks[key]
    )

    # Length score: 15 points.
    if 400 <= word_count <= 900:
        length_score = 15
    elif 250 <= word_count < 400 or 901 <= word_count <= 1100:
        length_score = 10
    elif 150 <= word_count < 250 or 1101 <= word_count <= 1300:
        length_score = 5
    else:
        length_score = 0

    # Technical skills score: 15 points.
    skill_count = len(skills)
    if skill_count >= 15:
        skills_score = 15
    elif skill_count >= 12:
        skills_score = 12
    elif skill_count >= 9:
        skills_score = 9
    elif skill_count >= 6:
        skills_score = 6
    elif skill_count >= 3:
        skills_score = 3
    else:
        skills_score = 0

    # Project detail score: 15 points.
    project_text = ""
    project_match = re.search(
        r"(?:projects|academic projects|personal projects)"
        r"(.*?)(?=\n\s*(?:experience|education|skills|certifications|"
        r"achievements|summary|objective)\b|$)",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if project_match:
        project_text = project_match.group(1)

    project_words = len(re.findall(r"\b[\w+#.-]+\b", project_text))
    if project_words >= 180:
        project_score = 15
    elif project_words >= 120:
        project_score = 12
    elif project_words >= 70:
        project_score = 8
    elif project_words >= 30:
        project_score = 4
    else:
        project_score = 0

    # Evidence of impact: 15 points.
    has_numbers = bool(
        re.search(r"\b\d+(?:\.\d+)?\s*(?:%|percent|users|customers|"
                  r"projects|days|hours|records|teams)\b", lower)
    )
    has_action_verbs = bool(
        re.search(
            r"\b(built|developed|designed|implemented|improved|"
            r"created|automated|optimized|achieved|led|reduced|"
            r"increased|delivered|deployed)\b",
            lower,
        )
    )

    impact_score = 0
    if has_action_verbs:
        impact_score += 7
    if has_numbers:
        impact_score += 8

    total = min(
        100,
        structural_score
        + length_score
        + skills_score
        + project_score
        + impact_score,
    )

    improvements = []

    if not sections["contact"]:
        improvements.append("Add professional contact information.")
    if not sections["summary"]:
        improvements.append("Add a concise professional summary or objective.")
    if not sections["skills"]:
        improvements.append("Add a clearly labelled technical skills section.")
    if not sections["projects"]:
        improvements.append("Include relevant projects and explain your contribution.")
    if not sections["experience"]:
        improvements.append(
            "Add internship or experience details; if you are a fresher, "
            "emphasize practical projects."
        )
    if not sections["education"]:
        improvements.append("Include your education and degree details.")
    if not sections["certifications"]:
        improvements.append("List relevant certifications if you have them.")
    if not sections["achievements"]:
        improvements.append("Add relevant achievements, awards, or competitions.")
    if not has_numbers:
        improvements.append(
            "Quantify project impact where accurate, using real metrics."
        )
    if word_count < 250:
        improvements.append("Add useful details; the resume may be too brief.")
    elif word_count > 1100:
        improvements.append("Review the resume for repetition and unnecessary detail.")
    if skill_count < 6:
        improvements.append("Include more relevant skills you can demonstrate.")

    summary = (
        "Strong resume structure"
        if total >= 80
        else "Good foundation with opportunities to improve"
        if total >= 60
        else "Several important areas could be strengthened"
    )

    return {
        "resume_score": total,
        "quality_score": total,
        "resume_quality": total,
        "resume_quality_score": total,
        "quality_checks": checks,
        "checks_passed": sum(checks.values()),
        "total_checks": len(checks),
        "detected_sections": sections,
        "sections_detected": sections,
        "improvements": improvements,
        "quality_summary": summary,
        "quality_breakdown": {
            "structure": structural_score,
            "length": length_score,
            "technical_skills": skills_score,
            "project_detail": project_score,
            "impact_evidence": impact_score,
        },
    }


# =========================================================
# 7. ROLE RECOMMENDATIONS
# =========================================================

def recommend_roles(resume_skills: list[str]) -> list[dict]:
    resume_set = {skill.lower() for skill in resume_skills}
    recommendations = []

    for role_name, required_skills in ROLE_SKILLS.items():
        matched = [
            skill for skill in required_skills
            if skill.lower() in resume_set
        ]
        missing = [
            skill for skill in required_skills
            if skill.lower() not in resume_set
        ]

        score = round(
            len(matched) / len(required_skills) * 100
        ) if required_skills else 0

        recommendations.append({
            "role": role_name,
            "title": role_name,
            "match_score": score,
            "matching_skills": matched,
            "missing_skills": missing,
        })

    recommendations.sort(
        key=lambda item: item["match_score"],
        reverse=True,
    )
    return recommendations


def prioritize_missing_skills(
    missing_skills: list[str],
) -> list[dict]:
    recommendations = []

    for skill in missing_skills:
        recommendations.append({
            "skill": skill,
            "priority": "High" if skill.lower() in {
                "python", "sql", "machine learning", "javascript"
            } else "Medium",
            "reason": f"Develop practical experience with {skill}.",
        })

    return recommendations


# =========================================================
# 8. JOB DESCRIPTION COMPARISON
# =========================================================

def compare_resume_to_job(
    resume_text: str,
    job_description: str,
) -> dict:
    resume_skills = extract_skills(resume_text)
    job_skills = extract_skills(job_description)

    resume_set = {skill.lower() for skill in resume_skills}
    job_set = {skill.lower() for skill in job_skills}

    matching_skills = sorted(resume_set.intersection(job_set))
    missing_skills = sorted(job_set.difference(resume_set))
    additional_skills = sorted(resume_set.difference(job_set))

    if job_set:
        match_score = round(
            len(matching_skills) / len(job_set) * 100
        )
        recommendation = (
            "Strong skill overlap"
            if match_score >= 75
            else "Moderate skill overlap"
            if match_score >= 45
            else "Consider developing the missing skills"
        )
    else:
        match_score = 0
        recommendation = (
            "No supported technical skills were detected in the job "
            "description. Review it manually for requirements."
        )

    return {
        "match_score": match_score,
        "resume_skills": resume_skills,
        "job_skills": job_skills,
        "matching_skills": matching_skills,
        "missing_skills": missing_skills,
        "additional_skills": additional_skills,
        "skill_recommendations": prioritize_missing_skills(missing_skills),
        "recommendation": recommendation,
        "match_method": "keyword_skill_overlap",
    }


# =========================================================
# 9. API ROOT AND HEALTH CHECK
# =========================================================

@app.get("/")
def root():
    return {
        "name": APP_NAME,
        "status": "running",
        "docs": "/docs",
        "endpoints": {
            "health": "/health",
            "analyze_resume": "/analyze",
            "match_jobs_legacy": "/match-jobs",
            "match_single_job": "/match-job",
            "supported_job_roles": "/jobs/roles",
            "search_jobs": "/jobs/search",
        },
    }


@app.get("/health")
def health():
    return {"status": "healthy"}


# =========================================================
# 10. RESUME ANALYSIS ENDPOINT
# =========================================================

@app.post("/analyze")
async def analyze_resume(
    resume: UploadFile = File(...),
    job_description: Optional[str] = Form(None),
):
    filename = resume.filename or "resume"
    extension = os.path.splitext(filename)[1].lower()

    if extension not in {".pdf", ".docx"}:
        raise HTTPException(
            status_code=400,
            detail="Upload a PDF or DOCX resume.",
        )

    file_bytes = await resume.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    if len(file_bytes) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="The maximum resume file size is 10 MB.",
        )

    resume_text = extract_resume_text(filename, file_bytes)

    if not resume_text:
        raise HTTPException(
            status_code=422,
            detail=(
                "No readable text was extracted. For scanned PDFs, "
                "check that Tesseract OCR is installed."
            ),
        )

    resume_skills = extract_skills(resume_text)
    word_count = len(re.findall(r"\b[\w+#.-]+\b", resume_text))

    quality = calculate_resume_quality(
        resume_text,
        resume_skills,
    )
    roles = recommend_roles(resume_skills)

    job_result = {
        "match_score": 0,
        "job_skills": [],
        "matching_skills": [],
        "missing_skills": [],
        "additional_skills": [],
        "skill_recommendations": [],
        "recommendation": "Add a job description to compare your resume with a role.",
    }

    if job_description and job_description.strip():
        job_result = compare_resume_to_job(
            resume_text,
            job_description,
        )

    return {
        "filename": filename,
        "word_count": word_count,
        "skills_count": len(resume_skills),
        "resume_skills": resume_skills,
        "skills": resume_skills,
        "recommended_roles": roles,
        "resume_text": resume_text,
        "extracted_text": resume_text,
        "job_skills": job_result["job_skills"],
        "matching_skills": job_result["matching_skills"],
        "missing_skills": job_result["missing_skills"],
        "additional_skills": job_result["additional_skills"],
        "skill_recommendations": job_result["skill_recommendations"],
        "match_score": job_result["match_score"],
        "recommendation": job_result["recommendation"],
        **quality,
    }


# =========================================================
# 11. LEGACY ENDPOINT FOR EXISTING FRONTENDS
# =========================================================

@app.post("/match-jobs")
async def match_jobs(
    resume: UploadFile = File(...),
    job_description: Optional[str] = Form(None),
):
    return await analyze_resume(resume, job_description)


# =========================================================
# 12. MATCH TEXT AGAINST A SINGLE JOB
# =========================================================

class JobMatchRequest(BaseModel):
    resume_text: str = Field(..., min_length=1)
    job_description: str = Field(..., min_length=1)


@app.post("/match-job")
def match_single_job(request: JobMatchRequest):
    if not request.resume_text.strip():
        raise HTTPException(
            status_code=400,
            detail="Resume text cannot be empty.",
        )

    if not request.job_description.strip():
        raise HTTPException(
            status_code=400,
            detail="Job description cannot be empty.",
        )

    return compare_resume_to_job(
        request.resume_text,
        request.job_description,
    )
